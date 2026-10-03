"""Reportes del administrador en Excel: semanal, cierre de cartera y clientes.

Dos zonas que nunca se mezclan:
- Norte: Ana (100.000 al 20 % = 120.000 en 4 cuotas de 30.000, paga una hoy)
  y Beto (prestamo ya pagado). La semana de Norte esta cuadrada.
- Sur: Carla (200.000 al 20 % = 240.000) y un prestamo anulado.
"""
import datetime
import io
import tempfile
from decimal import Decimal
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook


@pytest.fixture()
def entorno():
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker

    bd = Path(tempfile.gettempdir()) / "creditospro_reportes_admin_test.db"
    if bd.exists():
        bd.unlink()
    from app.database import (Base, Cliente, CuadreSemanal, Cuota, Empresa, Prestamo, Usuario,
                              Zona, hoy_local)
    from app.main import app
    from app.utils.company_activation import assign_company_key
    from app.utils.cuadre_semanal import calcular, lunes_de
    from app.utils.security import get_password_hash

    motor = create_engine(f"sqlite:///{bd}", connect_args={"check_same_thread": False})
    Base.metadata.create_all(motor)
    Sesion = sessionmaker(bind=motor, autoflush=False)

    def _sesion():
        db = Sesion()
        try:
            yield db
        finally:
            db.close()

    from conftest import sustituir_sesion, vaciar_limitador
    claves = sustituir_sesion(app, _sesion)
    hoy = hoy_local()
    lunes = lunes_de(hoy)
    d = {"hoy": hoy, "lunes": lunes}
    db = Sesion()
    try:
        e = Empresa(nombre="ReportesSA", activa=True)
        db.add(e); db.flush()
        d["empresa"] = e.id
        d["clave"] = assign_company_key(db, e)
        zonas = {}
        for cod, nombre in (("NO", "Norte"), ("SU", "Sur")):
            z = Zona(empresa_id=e.id, codigo=cod, nombre=nombre)
            db.add(z); db.flush()
            zonas[nombre] = z
        d["norte"], d["sur"] = zonas["Norte"].id, zonas["Sur"].id
        for u, rol in (("repjefa", "admin"), ("repcobra", "cobrador")):
            usr = Usuario(empresa_id=e.id, username=u, nombre=u.title(), rol=rol, activo=True,
                          password_hash=get_password_hash("ClaveDePrueba123!"))
            db.add(usr); db.flush()
            if rol == "cobrador":
                usr.zonas_asignadas.extend(zonas.values())

        def prestamo(nombre, cedula, zona, capital, estado="Activo", pagado=False):
            c = Cliente(empresa_id=e.id, cedula=cedula, nombre=nombre, telefono="3001112233",
                        zona_id=zona.id, activo=True)
            db.add(c); db.flush()
            total = capital * Decimal("1.2")
            p = Prestamo(empresa_id=e.id, cliente_id=c.id, zona_id=zona.id, capital=capital,
                         tasa_interes=Decimal("20"), total_pagar=total, num_cuotas=4,
                         valor_cuota=total / 4, plazo_dias=1, estado=estado,
                         fecha_inicio=hoy - datetime.timedelta(days=4), fecha_fin=hoy)
            db.add(p); db.flush()
            cuotas = []
            for n in range(4):
                cu = Cuota(empresa_id=e.id, prestamo_id=p.id, numero=n + 1, valor=total / 4,
                           valor_pagado=total / 4 if pagado else Decimal("0"),
                           estado="Pagada" if pagado else ("Anulada" if estado == "Anulado" else "Pendiente"),
                           fecha_vencimiento=hoy - datetime.timedelta(days=3 - n))
                db.add(cu); db.flush()
                cuotas.append(cu.id)
            return p.id, cuotas

        d["ana"], d["cuotas_ana"] = prestamo("Ana Norte", "2000001", zonas["Norte"], Decimal("100000"))
        d["beto"], _ = prestamo("Beto Norte", "2000002", zonas["Norte"], Decimal("50000"),
                                estado="Pagado", pagado=True)
        d["carla"], _ = prestamo("Carla Sur", "2000003", zonas["Sur"], Decimal("200000"))
        d["anulado"], _ = prestamo("Dora Sur", "2000004", zonas["Sur"], Decimal("80000"),
                                   estado="Anulado")
        db.commit()
    finally:
        db.close()

    def entrar(u):
        vaciar_limitador()
        cli = TestClient(app)
        cli.__enter__()
        assert cli.post("/license/activate", data={"license_key": d["clave"]}).status_code == 200
        assert cli.post("/auth/login", data={"username": u, "password": "ClaveDePrueba123!"}).status_code == 200
        cli.get("/ruta")
        cli.headers["x-csrf-token"] = cli.cookies.get("cp_csrf", "")
        return cli

    jefa, cobra = entrar("repjefa"), entrar("repcobra")
    # Ana paga una cuota hoy (en Norte); la semana de Norte queda cuadrada.
    r = cobra.post("/cobros/registrar", data={"cuota_id": d["cuotas_ana"][0], "valor_cobrado": "30000",
                                              "metodo_pago": "Efectivo"})
    assert r.status_code == 200, r.text
    db = Sesion()
    try:
        v = calcular(30000, 0, 5000, 0, 100000, 0, 125000, 5000)
        db.add(CuadreSemanal(empresa_id=d["empresa"], zona_id=d["norte"], semana=lunes,
                             verificado_por="Repjefa", **{k: v[k] for k in (
                                 "cobro", "prestamos", "gastos", "salarios", "base", "descuento",
                                 "efectivo", "esperado", "diferencia", "intereses", "utilidad")}))
        db.commit()
    finally:
        db.close()
    yield jefa, cobra, d, Sesion
    for c in (jefa, cobra):
        c.__exit__(None, None, None)
    for k in claves:
        app.dependency_overrides.pop(k, None)
    motor.dispose()
    bd.unlink(missing_ok=True)


def _libro(r):
    assert r.status_code == 200, r.text
    assert "spreadsheetml" in r.headers["content-type"]
    return load_workbook(io.BytesIO(r.content))


def _filas(ws):
    """Las filas de datos (debajo de la cabecera en la fila 5), como dicts."""
    cab = [c.value for c in ws[5]]
    out = []
    for fila in ws.iter_rows(min_row=6, values_only=True):
        if fila[0] in (None, ""):
            break
        out.append(dict(zip(cab, fila)))
    return out


def test_semanal_un_cuadre_por_zona_sin_mezclarlas(entorno):
    jefa, _, d, _ = entorno
    wb = _libro(jefa.get("/reportes/semanal", params={"semana": d["hoy"].isoformat()}))
    assert wb.sheetnames == ["Cuadre por zona", "Movimiento", "Cobros"]
    filas = {f["Zona"]: f for f in _filas(wb["Cuadre por zona"])}
    norte, sur, total = filas["Norte"], filas["Sur"], filas["TOTAL"]
    assert norte["Cobro"] == 30000 and norte["Efectivo"] == 125000 and norte["Esperado"] == 125000
    assert norte["Diferencia"] == 0 and norte["Estado"].startswith("Verificado por Repjefa")
    # Sur no se cuadro: lo que el sistema sabe, el resto en blanco.
    assert sur["Cobro"] == 0 and sur["Efectivo"] in (None, "") and sur["Estado"] == "Sin verificar"
    assert total["Cobro"] == 30000
    cobros = [f for f in _filas(wb["Cobros"]) if f["Cliente"]]
    assert [(c["Zona"], c["Cliente"], c["Valor"]) for c in cobros] == [("Norte", "Ana Norte", 30000)]
    mov = {f["Zona"]: f for f in _filas(wb["Movimiento"])}
    assert mov["Norte"]["Cobros"] == 1 and mov["Sur"]["Cobros"] == 0


def test_semanal_de_una_sola_zona(entorno):
    jefa, _, d, _ = entorno
    wb = _libro(jefa.get("/reportes/semanal", params={"zona_id": d["sur"]}))
    assert [f["Zona"] for f in _filas(wb["Cuadre por zona"])] == ["Sur", "TOTAL"]
    assert _filas(wb["Cobros"]) == [] or _filas(wb["Cobros"])[0]["Cliente"] != "Ana Norte"


def test_cierre_de_cartera_pide_el_inicio_de_los_ciclos(entorno):
    jefa, _, d, Sesion = entorno
    r = jefa.get("/reportes/cierre-cartera")
    assert r.status_code == 400 and "Finanzas" in r.json()["error"]
    from app.database import Empresa
    db = Sesion()
    try:
        db.get(Empresa, d["empresa"]).ciclo_inicio = d["lunes"] - datetime.timedelta(days=7)
        db.commit()
    finally:
        db.close()
    wb = _libro(jefa.get("/reportes/cierre-cartera"))
    assert wb.sheetnames == ["Resumen por zona", "Semana por semana", "Movimiento de cartera"]
    resumen = {f["Zona"]: f for f in _filas(wb["Resumen por zona"])}
    assert resumen["Norte"]["Cobrado"] == 30000 and resumen["Norte"]["Intereses"] == 5000
    assert resumen["Norte"]["Gastos"] == 5000 and resumen["Norte"]["Semanas cuadradas"] == "1/2"
    # Lo que queda por cobrar en Sur: Carla debe 240.000 = 200.000 capital + 40.000 interes
    # (el anulado no cuenta).
    assert resumen["Sur"]["Capital por cobrar"] == 200000
    assert resumen["Sur"]["Interés por cobrar"] == 40000
    # Dos semanas x dos zonas, cada semana con su total.
    semanas = [f for f in wb["Semana por semana"].iter_rows(min_row=6, values_only=True) if f[1]]
    assert sum(1 for f in semanas if f[1] in ("Norte", "Sur")) == 4
    assert jefa.get("/reportes/cierre-cartera", params={"ciclo": 9}).status_code == 400


def test_clientes_una_fila_por_prestamo(entorno):
    jefa, _, d, _ = entorno
    filas = _filas(_libro(jefa.get("/reportes/clientes-prestamos"))["Clientes y préstamos"])
    por_cliente = {f["Cliente"]: f for f in filas if f["Préstamo"]}
    # Con saldo: Ana y Carla; ni el pagado de Beto ni el anulado.
    assert set(por_cliente) == {"Ana Norte", "Carla Sur"}
    ana = por_cliente["Ana Norte"]
    assert (ana["Total"], ana["Pagado"], ana["Saldo"]) == (120000, 30000, 90000)
    assert (ana["Saldo capital"], ana["Saldo interés"]) == (75000, 15000)
    assert ana["Cuotas pagadas"] == 1 and ana["Cuotas atrasadas"] == 2 and ana["Semáforo"] == "Amarillo"
    todos = _filas(_libro(jefa.get("/reportes/clientes-prestamos",
                                   params={"estado": "todos"}))["Clientes y préstamos"])
    assert sum(1 for f in todos if isinstance(f["Préstamo"], int)) == 4
    sur = _filas(_libro(jefa.get("/reportes/clientes-prestamos",
                                 params={"zona_id": d["sur"], "estado": "anulados"}))["Clientes y préstamos"])
    assert [f["Cliente"] for f in sur if isinstance(f["Préstamo"], int)] == ["Dora Sur"]
    assert jefa.get("/reportes/clientes-prestamos", params={"estado": "x"}).status_code == 400


def test_los_reportes_son_del_admin(entorno):
    _, cobra, _, _ = entorno
    for url in ("/reportes/semanal", "/reportes/cierre-cartera", "/reportes/clientes-prestamos"):
        assert cobra.get(url).status_code == 403, url


def test_la_pagina_ofrece_los_tres_reportes(entorno):
    jefa, _, d, _ = entorno
    html = jefa.get("/reportes").text
    for pieza in ("bajar('semanal'", "bajar('clientes-prestamos'", 'id="rs-semana"',
                  d["lunes"].isoformat(), "Primero indica el lunes"):
        assert pieza in html, pieza
